"""Dump column headers of IRS SOI tables 6.1 and 3.2 (Form 1120-S) to design the parser."""
import io
import sys

import openpyxl

sys.path.insert(0, "etl")
from netutil import urlopen  # noqa: E402


def load(name):
    with urlopen("https://www.irs.gov/pub/irs-soi/" + name, timeout=120) as r:
        return openpyxl.load_workbook(io.BytesIO(r.read()), read_only=True, data_only=True).worksheets[0]


for name in ["22co61ccr.xlsx", "19co61ccr.xlsx"]:
    rows = list(load(name).iter_rows(values_only=True))
    print(f"===== {name}: {len(rows)} rows x {max(len(r) for r in rows)} cols")
    for i, r in enumerate(rows[:14]):
        print(f"  row{i}: first={str(r[0])[:40]!r} nonempty={sum(1 for c in r if c not in (None, ''))}")
    first_data = next(i for i, r in enumerate(rows) if r[0] and str(r[0]).strip().startswith("Number of returns"))
    print("  first data row:", first_data)
    ncols = max(len(r) for r in rows)
    for c in range(1, ncols):
        head = " / ".join(str(rows[i][c]).replace("\n", " ").strip() for i in range(first_data)
                          if c < len(rows[i]) and rows[i][c] not in (None, ""))
        vals = {str(r[0]).strip(): r[c] for r in rows[first_data:] if r[0] and c < len(r)}
        print(f"  col{c}: {head[:150]!r} receipts={vals.get('Business receipts')} netinc={vals.get('Net income (less deficit)')}")
        if c > 60 and name.startswith("19"):
            break

rows = list(load("22co32ccr.xlsx").iter_rows(values_only=True))
print("===== 22co32ccr.xlsx")
for i, r in enumerate(rows[:10]):
    print(f"  row{i}:", [str(c)[:30] for c in r if c not in (None, "")])
for r in rows[10:120]:
    if r[0]:
        print("  ", str(r[0]).strip()[:50], "|", [c for c in r[1:14]])
